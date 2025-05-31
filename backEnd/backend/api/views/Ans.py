from datetime import date
from google.generativeai import types
import google.generativeai as genai
from datetime import datetime, time, timedelta
from google.genai import types
from google.generativeai.types import FunctionDeclaration
from google.generativeai.types import  FunctionDeclaration
from ..utils.Neo4jSp import get_key_phase
from google.generativeai import protos
# import google.generativeai as genai
from rest_framework.decorators import api_view
from django.views.decorators.csrf import csrf_exempt
from django.http import JsonResponse
from ..utils.neo4j_driver import driver
from ..utils.Neo4jSp import run_query
from ..models.component import component
import re
from docx import Document
import string
from ..models.gConc import GConc
from ..models.gRel import gRel
import os
from django.conf import settings
from pathlib import Path


countkey = 0
countuse = 0
dateAns = datetime.now()
list_key = ["AIzaSyDyVhyjt02Alp9YiMArPKPx6XC5RVkAJ18"]
limitkey = len(list_key)
limituse = 9

def component_in_node4j(listtrip: list[tuple[str,str,str]]):
    # print(listtrip)
    return listtrip


def call_genai(Text):
    # return [("Tổ chức","ngoài","khu công nghiệp")]
    global countkey, countuse, dateAns
    if is_now_in_custom_range(date_obj=dateAns):
        if countuse < limituse:
            countuse += 1
        else:
            countuse = 0
            countkey +=1
    else:
        countkey = 0
        countuse = 1
        dateAns = date()


    genai.configure(api_key=list_key[countkey])
    model = genai.GenerativeModel("gemini-2.0-flash")

    try:
        functions = [
    {
        "name": "component_in_node4j",
        "description": "Tập các bộ ba tri thức",
        "parameters": {
            "type": "object",
            "properties": {
                "listtrip": {
                    "type": "array",
                    "description": "Danh sách các tuple, mỗi tuple chứa 3 chuỗi đại diện cho bộ ba tri thức",
                    "items": {
                        "type": "array",
                        "description": "Tuple chứa 3 chuỗi",
                        "items": {
                            "type": "string",
                            "description": "Thành phần của bộ ba tri thức (entity1, relation, entity2)"
                        },
                        "minItems": 3,
                        "maxItems": 3
                    }
                }
            },
            "required": ["listtrip"]
        }
    }
]
        
        function_declarations = []
        for func in functions:
            description = (
            f"{func['description']}. Hàm trả về một list chứa các tuple, "
            "mỗi tuple chứa đúng 3 chuỗi đại diện cho bộ ba tri thức (entity1, relation, entity2)."
        )
            properties = {
            "listtrip": genai.protos.Schema(
                type=genai.protos.Type.ARRAY,
                description=func["parameters"]["properties"]["listtrip"]["description"],
                items=genai.protos.Schema(
                    type=genai.protos.Type.ARRAY,
                    description=func["parameters"]["properties"]["listtrip"]["items"]["description"],
                    items=genai.protos.Schema(
                        type=genai.protos.Type.STRING,
                        description=func["parameters"]["properties"]["listtrip"]["items"]["items"]["description"]
                    )
                )
            )
        }

            func_decl = genai.protos.FunctionDeclaration(
            name=func["name"],
            description=description,
            parameters=genai.protos.Schema(
                type=genai.protos.Type.OBJECT,
                properties=properties,
                required=func["parameters"].get("required", [])
            )
        )
            function_declarations.append(func_decl) 

        
        # tools = types.Tool(function_declarations=function_declarations)
        


        contents = { "text":f"""Bạn sẽ nhận một đoạn văn bản và 2 dictionary, mỗi dictionary có 1 key .
                                                        dictionary 1 (concept): danh sách các khái niệm, cách nhau bởi dấu phẩy (,)
                                                        dictionary 2 (relation): danh sách các cụm quan hệ, cách nhau bởi dấu phẩy (,)
                                                    Nhiệm vụ của bạn là:
                                                        Phân tích đoạn văn bản đầu vào thành các bộ ba kiến thức (A = B => C).
                                                        Trong đó:
                                                            A và C là các khái niệm, được chọn từ dictionary Concept. Nếu không tìm thấy khớp chính xác, hãy tìm từ đồng nghĩa gần nhất trong danh sách. Nếu vẫn không có, thay bằng dấu *.
                                                            B là quan hệ (relation), Nếu không khớp hoàn toàn, hãy tìm cụm từ đồng nghĩa gần nhất. Nếu không tìm được thì không ghi nhận bộ ba tri thức.
                                                    Đầu ra yêu cầu là một danh sách các bộ ba có dạng A = B => C, ví dụ:
                                                        “Hợp đồng lao động = bị chấm dứt bởi => Người sử dụng lao động”
                                                        “* = được miễn => Thuế thu nhập cá nhân”
                                                        "Đối tượng = áp dụng =>*"
                                                    Hãy chỉ in ra danh sách các bộ ba dạng A = B => C. Không cần giải thích thêm.
                                                    Đoạn văn bản đầu vào là {Text}
                                                    dictionary 1: {get_key_phase("listconcept.docx")}
                                                    dictionary 1: {get_key_phase("listrelation.docx")}"""}
        print(get_key_phase("listconcept.docx"))
        print("\n\n\n\n")
        print(get_key_phase("listrelation.docx"))
        
#         pdf_path = os.path.join(settings.MEDIA_ROOT, "merge_content.pdf")
#         filepath = Path(pdf_path)
#         if not filepath.is_file():
#             raise FileNotFoundError(f"Không tìm thấy tệp: {pdf_path}")
#         pdf_blob = {
#     "mime_type": "application/pdf",
#     "data": filepath.read_bytes()
# }
         
        # print(contents)
    # Send request with function declarations
        response = model.generate_content(
            #   contents=[pdf_blob,contents], tools=[genai.protos.Tool(function_declarations=function_declarations)]
              contents=[contents], tools=[genai.protos.Tool(function_declarations=function_declarations)]
        )

        tool_call = response.candidates[0].content.parts[0].function_call
        
        if tool_call.name == "component_in_node4j":
            # print(tool_call.args)
            result = component_in_node4j(**tool_call.args)
            # print(f"Function execution result: {result}")
            # return result
        # print("gần rồi", result)
        return result
        
    except Exception as e:
        print(e)
        return [("Result","trả về","None")]
    
def get_id_trip(trip: list[tuple[str,str,str]]):
    result = set()
    temp_result = set()
    for item in trip:
        if item[0] == '*' and item[1] == '*' and item[2] == '*':
            where = ''
        else:
            where = f"""WHERE {' ' if item[0] == '*' else f"'{item[0].lower()}' in a.lstConC"}\
{' ' if item[1] == '*' else f" AND '{item[1].lower()}' in r.lstRel"}\
{' ' if item[2] == '*' else f" AND '{item[2].lower()}' in b.lstConC"}"""

        query = f"""
            MATCH (a)-[r]-(b)
            {where}
            RETURN a.descendants AS d_S, r.descendants AS d_R, b.descendants AS d_O
            """
        list_des = run_query(driver=driver, query=query)
        if not list_des:
            continue
        for item in list_des:
            temp_result = intersection(item["d_S"],item["d_R"])
            temp_result = intersection(temp_result,item["d_O"])
        if len(result) == 0:
            result = temp_result
        else:
            result = result & temp_result
    return result

def get_law(set_id: set):
    list_id = list(set_id)
    list_id.sort()
    for item in list_id:
        id = component.objects.get(id=item)
        id.descendants.reverse()
    
def intersection(list1:list, list2:list):
    list1 = set(list1)
    list2 = set(list2)
    return list1 & list2

def get_all_parents_chain(item_id):
    """
    Truy ngược lên các phần tử cha theo chuỗi cha-con.
    """
    result = []
    current_id = item_id

    while True:
        try:
            component1 = component.objects.get(id=current_id)
            result.insert(0, (component1.name,component1.contentSearch))
        except component1.DoesNotExist:
            break

        parents = component1.parent or -1
        if not parents:
            continue

        current_id = parents

    return result


def get_all_children_tree_order(root_id):
    """
    Duyệt cây con theo DFS từ trái sang phải, không dùng đệ quy.
    """
    result = []
    stack = [root_id]

    while stack:
        current_id = stack.pop()

        try:
            component1 = component.objects.get(id=current_id)
            result.append((component1.name, component1.contentSearch))
        except component1.DoesNotExist:
            continue

        children = component1.child or []
        for child_id in reversed(children):  # reversed để đúng thứ tự trái → phải
            stack.append(child_id)

    return result[1:]

def Ans_theory(text):
    theory = get_theory(text=text)
    result1= gRel.objects.filter(lstRel__contains = theory).first()
    result2= GConc.objects.filter(lstConC__contains = theory).first()
    # print("text","1", result2.meaning,"2",result1)
    print("text2",bool(result1 and result2))
    if not result1 and not result2:
        return False
    if result1:
        result1 = result1.meaning
    if result2:
        result2 = result2.meaning
    result1 = result2 if result2 else result1
    print("result", result1)
    result1 = component.objects.filter(contentSearch__icontains = result1).first()
    if result1:
        
        return [{"content": get_result(result1.id), "idlaw": result1.idLaw_id }]
    else:
        return False

def Ans_action(text):
    trip = call_genai(text)
    print(trip)
    trip = get_id_trip(trip=trip)
    print(trip)
    trip = filter_top_level_components(trip)
    result = list()
    if not trip:
        return False
    else:
        # trip = list(trip).sort()
        for item in trip:
            value = get_result(item)
            if value:
                result.append({"content": value, "idlaw": component.objects.filter(id=item).first().idLaw_id })
            
    return result

@csrf_exempt
@api_view(['POST'])
def Ans(request):
    text = request.data.get("text")
    if not text:
        return JsonResponse({"content": "Xin lỗi về sự bất tiện\n Hệ thống chưa nhận được câu hỏi"}, status = 200)
        
    theory = Ans_theory(text=text)
    print(theory)
    if theory:
        return JsonResponse(theory,safe=False,status = 200)
    action = Ans_action(text=text)
    if action:
        return JsonResponse(action,safe=False,status = 200)
    if theory and action:
        return JsonResponse({"content": "Xin lỗi về sự bất tiện\n Hệ thống chưa được học về tri được tra cứu"}, status = 200)
    return JsonResponse({"content": " 1Xin lỗi về sự bất tiện\n Hệ thống chưa được học về tri được tra cứu"}, status = 200)
    
def extract_question_content(question):
    question = question.lower().strip()

    patterns = [
        r'\b(là gì|thế nào là|được hiểu là gì|được định nghĩa như thế nào|định nghĩa|khái niệm về|hiểu như thế nào là|ý nghĩa của|được gọi là gì|nghĩa là gì|có nghĩa là gì|có thể hiểu là gì|được coi là gì|là khái niệm chỉ|là thuật ngữ dùng để|là một loại|được gọi là|có thể được hiểu là|dùng để chỉ|được định nghĩa là|có thể được coi là|được xem là|mang nghĩa là)\b'
    ]

    for pattern in patterns:
        question = re.sub(pattern, '', question)

    question = question.translate(str.maketrans('', '', string.punctuation))

    return question.strip()


def get_concepts_and_relations():
    doc1 = Document(os.path.join(settings.MEDIA_ROOT, "listconcept.docx"))
    doc2 = Document(os.path.join(settings.MEDIA_ROOT, "listrelation.docx"))
    tables1 = doc1.tables
    tables2 = doc2.tables

    concept_line = tables1[0].rows[1].cells[0].text
    concept_list = [c.strip().lower() for c in concept_line.split(",")]

    relation_line = tables2[0].rows[1].cells[0].text
    relation_list = [r.strip().lower() for r in relation_line.split(",")]

    return concept_list, relation_list

def check_text(text, listtheory: tuple):
    text = text.lower().strip()

    for concept in listtheory[0]:
        if text == concept.lower():
            return concept

    for relation in listtheory[1]:
        if text == relation.lower():
            return relation

    return False

def get_theory(text):
    theory = extract_question_content(text)
    theory = check_text(theory, get_concepts_and_relations())
    if not theory:
        return False
    else:
        return theory

def get_result(id_result):
    parents = get_all_parents_chain(id_result)
    descendants = get_all_children_tree_order(id_result)
    header = list()
    content = ''
    if parents:
        for item in parents:
            header.append(item[0])
            content = f"{content}\n{item[0]}{'' if not item[1] else f': {item[1]}'}"
    if descendants: 
        for item in descendants:
            content = f"{content}\n{item[0]}{'' if not item[1] else f': {item[1]}'}"

    return f'Theo {header}\n {content} \n\n' if header else ""


# def call_genai(Text):
#     # return [("Tổ chức","ngoài","khu công nghiệp")]
#     # genai.configure(api_key="AIzaSyA96dLgKKqz8zXx_YcZJCyFYkC5jBoVAvM")
#     client = genai.Client(api_key="AIzaSyA96dLgKKqz8zXx_YcZJCyFYkC5jBoVAvM")

#     try:
#         functions = [
#     {
#         "name": "component_in_node4j",
#         "description": "Tập các bộ ba tri thức",
#         "parameters": {
#             "type": "object",
#             "properties": {
#                 "listtrip": {
#                     "type": "array",
#                     "description": "Danh sách các tuple, mỗi tuple chứa 3 chuỗi đại diện cho bộ ba tri thức",
#                     "items": {
#                         "type": "array",
#                         "description": "Tuple chứa 3 chuỗi",
#                         "items": {
#                             "type": "string",
#                             "description": "Thành phần của bộ ba tri thức (entity1, relation, entity2)"
#                         },
#                         "minItems": 3,
#                         "maxItems": 3
#                     }
#                 }
#             },
#             "required": ["listtrip"]
#         }
#     }
# ]
         
#         # model = genai.GenerativeModel("gemini-2.0-flash")

        
#         # tools = types.Tool(function_declarations=function_declarations)
        


#         contents =  { "text":f"Bạn sẽ nhận một đoạn văn bản và file PDF gồm 2 bảng, mỗi bảng có 2 dòng và 1 cột:.\n \
#                                                         Bảng 1 (concept):\n \
#                                                                 Dòng 1: tiêu đề là “Concepts”\n\
#                                                                 Dòng 2: danh sách các khái niệm, cách nhau bởi dấu phẩy (,)\n\
#                                                         Bảng 2 (relation).\n\
#                                                                 Dòng 1: tiêu đề là “Relations”\n\
#                                                                 Dòng 2: danh sách các cụm quan hệ, cách nhau bởi dấu phẩy (,)\n\
#                                                     Nhiệm vụ của bạn là:\n\
#                                                         Phân tích đoạn văn bản đầu vào thành các bộ ba kiến thức (A = B => C).\n\
#                                                         Trong đó:\n\
#                                                             A và C là các khái niệm, được chọn từ bảng Concept. Nếu không tìm thấy khớp chính xác, hãy tìm từ đồng nghĩa gần nhất trong danh sách. Nếu vẫn không có, thay bằng dấu *.\n\
#                                                             B là quan hệ (relation), Nếu không khớp hoàn toàn, hãy tìm cụm từ đồng nghĩa gần nhất. Nếu không tìm được.\n\
#                                                     Đầu ra yêu cầu là một danh sách các bộ ba có dạng A = B => C, ví dụ:\n\
#                                                         “Hợp đồng lao động = bị chấm dứt bởi => Người sử dụng lao động”\n\
#                                                         “* = được miễn => Thuế thu nhập cá nhân”\n\
#                                                     Hãy chỉ in ra danh sách các bộ ba dạng A = B => C. Không cần giải thích thêm.\n\
#                                                     Đoạn văn bản đầu vào là {Text}"}
        
#         pdf_path = os.path.join(settings.MEDIA_ROOT, "merge_content.pdf")
#         filepath = Path(pdf_path)
#         if not filepath.is_file():
#             raise FileNotFoundError(f"Không tìm thấy tệp: {pdf_path}")
#         pdf_blob = client.files.upload(file=filepath, config=dict(
#     mime_type='application/pdf'))        

#     # Send request with function declarations
#         response = model.generate_content(
#               contents=[pdf_blob,contents], tools=[genai.protos.Tool(function_declarations=function_declarations)]
#         )

#         tool_call = response.candidates[0].content.parts[0].function_call
        
#         if tool_call.name == "component_in_node4j":
#             # print(tool_call.args)
#             result = component_in_node4j(**tool_call.args)
#             print(f"Function execution result: {result}")
#             # return result
#         # print("gần rồi", result)
#         return result
        
#     except Exception as e:
#         print(e)
#         return [("Result","trả về","None")]


def is_now_in_custom_range(date_obj):
    now = datetime.now()
    
    start = datetime.combine(date_obj, time(7, 0, 0))  # 07:00 sáng ngày đó
    end = start + timedelta(days=1) - timedelta(seconds=1)  # 06:59:59 sáng hôm sau

    return start <= now <= end


def filter_top_level_components(ids_input):
    components = component.objects.filter(id__in=ids_input)
    id_to_descendants = {c.id: c.descendants or [] for c in components}

    result = set(ids_input)

    for id1 in ids_input:
        for id2 in ids_input:
            if id1 != id2:
                if id1 in id_to_descendants.get(id2, []):
                    result.discard(id1)

    return list(result)

